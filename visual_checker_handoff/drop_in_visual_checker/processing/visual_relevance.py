from __future__ import annotations

import base64
import hashlib
import io
import json
import re
import threading
import time
from concurrent.futures import FIRST_COMPLETED, ThreadPoolExecutor, wait
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable
from urllib.parse import urljoin

import openpyxl
import requests
from bs4 import BeautifulSoup
from openai import OpenAI


DEFAULT_MODEL = "gpt-4.1-mini"
DEFAULT_REQUEST_TIMEOUT = 15
DEFAULT_STRAGGLER_TIMEOUT = 90
DEFAULT_RETRY_BATCH_SIZE = 25
DEFAULT_MAX_IMAGES_PER_URL = 4
TRANSIENT_FAILURE_VERDICTS = {"ERROR", "FETCH ERROR", "RATE LIMITED"}
BLOCKED_FAILURE_VERDICTS = {"BLOCKED", "NOT FOUND"}
MODEL_PRICING_PER_1M = {
    "gpt-4.1-mini": {
        "input": 0.40,
        "cached_input": 0.10,
        "output": 1.60,
    }
}
USER_AGENT = (
    "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/126.0 Safari/537.36"
)
JUNK_IMAGE = re.compile(
    r"logo|sprite|icon|avatar|placeholder|1x1|pixel|spacer|blank\.|"
    r"/ads?/|doubleclick|/badge|favicon|amazon-adsystem|gravatar|tracking",
    re.I,
)
PROMPT_TEMPLATE = """You are auditing media coverage images.

Image relevance criteria:
{criteria}

Look at this image and decide whether it matches the relevance criteria.

Answer NO if it is a stock photo, headshot, product-only shot, logo, press-release banner, unrelated news photo, or generic decoration unless the criteria explicitly say otherwise.

Caption or alt text found with the image: {caption}

Reply with exactly one line:
YES | five word reason
or
NO | five word reason"""


@dataclass(frozen=True)
class ImageCandidate:
    url: str
    caption: str
    source: str
    score: int


def workbook_key(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def normalize_url(value: Any) -> str:
    if value is None:
        return ""
    url = str(value).strip()
    if not url or url.lower() == "nan":
        return ""
    if url.startswith(("http://", "https://")):
        return url
    if "." in url and " " not in url:
        return f"https://{url}"
    return url


def load_cache(cache_path: str | Path) -> dict[str, dict[str, str]]:
    path = Path(cache_path)
    if not path.exists():
        return {}
    try:
        return normalize_cache_labels(json.loads(path.read_text()))
    except (json.JSONDecodeError, OSError):
        return {}


def save_cache(cache: dict[str, dict[str, str]], cache_path: str | Path) -> None:
    path = Path(cache_path)
    tmp_path = path.with_suffix(".tmp")
    tmp_path.write_text(json.dumps(cache, indent=1, sort_keys=True))
    tmp_path.replace(path)


def get_sheet_data(workbook_bytes: bytes, sheet_name: str) -> tuple[list[str], list[list[Any]]]:
    workbook = openpyxl.load_workbook(io.BytesIO(workbook_bytes), read_only=True, data_only=True)
    worksheet = workbook[sheet_name]
    rows = list(worksheet.iter_rows(values_only=True))
    if not rows:
        return [], []
    headers = [
        str(value).strip() if value is not None else f"Column {index}"
        for index, value in enumerate(rows[0], start=1)
    ]
    return headers, [list(row) for row in rows[1:]]


def likely_url_columns(headers: list[str]) -> list[str]:
    preferred = [name for name in headers if "url" in name.lower() or "link" in name.lower()]
    return preferred + [name for name in headers if name not in preferred]


def urls_from_rows(rows: list[list[Any]], headers: list[str], url_column: str) -> list[tuple[int, str]]:
    if url_column not in headers:
        return []
    url_column_index = headers.index(url_column)
    urls_by_row = []
    for index, row in enumerate(rows):
        value = row[url_column_index] if url_column_index < len(row) else ""
        url = normalize_url(value)
        if url.startswith(("http://", "https://")):
            urls_by_row.append((index, url))
    return urls_by_row


def fetch_page(url: str, session: requests.Session, request_timeout: int) -> BeautifulSoup:
    response = session.get(url, timeout=request_timeout, headers={"User-Agent": USER_AGENT})
    response.raise_for_status()
    return BeautifulSoup(response.text, "html.parser")


def meta_content(soup: BeautifulSoup, name: str) -> str:
    tag = soup.find("meta", property=name) or soup.find("meta", attrs={"name": name})
    if not tag:
        return ""
    return str(tag.get("content") or "").strip()


def collect_image_candidates(url: str, soup: BeautifulSoup) -> list[ImageCandidate]:
    candidates: list[ImageCandidate] = []
    for index, prop in enumerate(("og:image", "og:image:secure_url", "twitter:image")):
        src = meta_content(soup, prop)
        if src and not JUNK_IMAGE.search(src):
            caption = meta_content(soup, "og:image:alt") or meta_content(soup, "twitter:image:alt")
            candidates.append(ImageCandidate(urljoin(url, src), caption, prop, 1_000_000 - index))

    for img in soup.find_all("img"):
        src = img.get("src") or img.get("data-src") or img.get("data-lazy-src") or img.get("data-original")
        if not src or JUNK_IMAGE.search(str(src)):
            continue
        width = parse_int(img.get("width"))
        height = parse_int(img.get("height"))
        fig = img.find_parent("figure")
        caption = ""
        if fig and fig.find("figcaption"):
            caption = fig.find("figcaption").get_text(" ", strip=True)
        caption = caption or str(img.get("alt") or "").strip()
        candidates.append(ImageCandidate(urljoin(url, str(src)), caption, "inline", width * height if width and height else 1))

    return sorted(candidates, key=lambda candidate: candidate.score, reverse=True)


def fetch_candidate_image(candidate: ImageCandidate, session: requests.Session, request_timeout: int) -> tuple[bytes, str, str]:
    try:
        response = session.get(candidate.url, timeout=request_timeout, headers={"User-Agent": USER_AGENT})
        response.raise_for_status()
    except requests.RequestException as exc:
        return b"", "", f"{type(exc).__name__}: {exc}"[:160]
    media_type = response.headers.get("Content-Type", "image/jpeg").split(";")[0].strip()
    if media_type not in {"image/jpeg", "image/png", "image/gif", "image/webp"}:
        return b"", "", f"unsupported type {media_type}"
    if len(response.content) < 4000:
        return b"", "", "image too small to be editorial"
    return response.content, media_type, ""


def classify_image(client: OpenAI, model: str, image: bytes, media_type: str, caption: str, criteria: str) -> tuple[str, str, dict[str, int]]:
    image_data = base64.b64encode(image).decode("ascii")
    response = client.responses.create(
        model=model,
        max_output_tokens=60,
        input=[
            {
                "role": "user",
                "content": [
                    {"type": "input_image", "image_url": f"data:{media_type};base64,{image_data}"},
                    {
                        "type": "input_text",
                        "text": PROMPT_TEMPLATE.format(
                            criteria=(criteria or "Use the user's stated relevance criteria.").strip(),
                            caption=(caption[:500] or "none"),
                        ),
                    },
                ],
            }
        ],
    )
    line = (response.output_text or "").strip().splitlines()[0] if response.output_text else ""
    verdict, _, reason = line.partition("|")
    verdict = verdict.strip().upper()
    return "YES" if verdict.startswith("YES") else "NO", reason.strip() or "model returned no reason", response_usage(response)


def classify_candidate_images(client: OpenAI, model: str, candidates: list[ImageCandidate], session: requests.Session, request_timeout: int, max_images_per_url: int, criteria: str) -> dict[str, str]:
    usage_total = zero_usage()
    checked = 0
    skipped = 0
    no_results: list[tuple[ImageCandidate, str]] = []
    last_skip_reason = ""

    for candidate in dedupe_candidates(candidates)[:max_images_per_url]:
        image, media_type, skip_reason = fetch_candidate_image(candidate, session, request_timeout)
        if skip_reason:
            skipped += 1
            last_skip_reason = skip_reason
            continue
        checked += 1
        verdict, reason, usage = classify_image(client, model, image, media_type, candidate.caption, criteria)
        usage_total = add_usage(usage_total, usage)
        if verdict == "YES":
            return image_result(candidate, verdict, reason, usage_total, checked, skipped)
        no_results.append((candidate, reason))

    if no_results:
        candidate, reason = no_results[0]
        if len(no_results) > 1:
            reason = f"{reason}; checked {len(no_results)} images"
        return image_result(candidate, "NO", reason, usage_total, checked, skipped)

    return {
        "verdict": "NO IMAGE",
        "image_url": "",
        "reason": last_skip_reason or "no supported editorial image found",
        "source": "",
        "images_checked": str(checked),
        "images_skipped": str(skipped),
        "input_tokens": "0",
        "cached_input_tokens": "0",
        "output_tokens": "0",
    }


def process_url(url: str, api_key: str, model: str, request_timeout: int, max_images_per_url: int, criteria: str) -> dict[str, str]:
    session = requests.Session()
    client = OpenAI(api_key=api_key, timeout=request_timeout)
    try:
        soup = fetch_page(url, session, request_timeout)
        candidates = collect_image_candidates(url, soup)
        if not candidates:
            return empty_result("NO IMAGE", "no candidate images found")
        return classify_candidate_images(client, model, candidates, session, request_timeout, max_images_per_url, criteria)
    except Exception as exc:
        verdict, reason = label_fetch_failure(exc)
        return empty_result(verdict, reason[:240])


def process_batch(
    urls_by_row: list[tuple[int, str]],
    api_key: str,
    model: str,
    batch_size: int,
    workers: int,
    request_timeout: int,
    straggler_timeout: int,
    max_images_per_url: int,
    criteria: str,
    cache_path: str | Path = "visual_cache.json",
    retry_errors: bool = False,
    retry_verdicts: set[str] | None = None,
    progress_callback: Callable[[int, int, int], None] | None = None,
) -> tuple[int, int, int, dict[str, dict[str, str]], dict[str, int]]:
    cache = load_cache(cache_path)
    if retry_verdicts:
        retry_kind = "negative"
        pending = retry_pending(urls_by_row, cache, retry_verdicts, retry_kind, batch_size)
    elif retry_errors:
        retry_kind = "error"
        pending = retry_pending(urls_by_row, cache, TRANSIENT_FAILURE_VERDICTS, retry_kind, batch_size)
    else:
        retry_kind = None
        pending = [(row, url) for row, url in urls_by_row if url not in cache][:batch_size]
    pending = dedupe_pending(pending)
    if not pending:
        return 0, 0, 0, cache, zero_usage()

    cache_lock = threading.Lock()
    completed = 0
    timed_out = 0
    batch_usage = zero_usage()

    for chunk in chunked(pending, max(1, workers)):
        executor = ThreadPoolExecutor(max_workers=max(1, workers))
        futures = {
            executor.submit(process_url, url, api_key, model, request_timeout, max_images_per_url, criteria): (row, url)
            for row, url in chunk
        }
        remaining = set(futures)
        deadline = time.monotonic() + straggler_timeout
        try:
            while remaining and time.monotonic() < deadline:
                done, remaining = wait(remaining, timeout=0.5, return_when=FIRST_COMPLETED)
                for future in done:
                    _, url = futures[future]
                    previous = cache.get(url, {})
                    try:
                        result = future.result()
                    except Exception as exc:
                        verdict, reason = label_fetch_failure(exc)
                        result = empty_result(verdict, reason[:240])
                    result = with_retry_metadata(result, previous, retry_kind)
                    with cache_lock:
                        cache[url] = result
                        batch_usage = add_usage(batch_usage, result_usage(result))
                        completed += 1
                        if completed % 5 == 0:
                            save_cache(cache, cache_path)
                        if progress_callback:
                            progress_callback(completed, timed_out, len(pending))
            if remaining:
                timed_out += len(remaining)
                for future in remaining:
                    future.cancel()
                if progress_callback:
                    progress_callback(completed, timed_out, len(pending))
        finally:
            executor.shutdown(wait=False, cancel_futures=True)

    save_cache(cache, cache_path)
    return len(pending), completed, timed_out, cache, batch_usage


def workbook_bytes_with_results(source_bytes: bytes, sheet_name: str, url_column_name: str, results: dict[int, dict[str, str]]) -> bytes:
    workbook = openpyxl.load_workbook(io.BytesIO(source_bytes))
    worksheet = workbook[sheet_name]
    headers = [worksheet.cell(1, col).value for col in range(1, worksheet.max_column + 1)]
    try:
        url_col = headers.index(url_column_name) + 1
    except ValueError as exc:
        raise ValueError(f"URL column '{url_column_name}' was not found") from exc

    output_columns = {
        "Has Relevant Visuals?": None,
        "Visual Checked URL": None,
        "Visual Check Reason": None,
        "Visual Check Source": None,
        "Visual Images Checked": None,
    }
    for col in range(1, worksheet.max_column + 1):
        value = worksheet.cell(1, col).value
        if value in output_columns:
            output_columns[value] = col
    next_col = worksheet.max_column + 1
    for title, col in output_columns.items():
        if col is None:
            output_columns[title] = next_col
            worksheet.cell(1, next_col, title)
            next_col += 1

    for row_index, result in results.items():
        row = row_index + 2
        url = normalize_url(worksheet.cell(row, url_col).value)
        if not url:
            continue
        worksheet.cell(row, output_columns["Has Relevant Visuals?"], result["verdict"])
        worksheet.cell(row, output_columns["Visual Checked URL"], result["image_url"])
        worksheet.cell(row, output_columns["Visual Check Reason"], result["reason"][:900])
        worksheet.cell(row, output_columns["Visual Check Source"], result["source"])
        worksheet.cell(row, output_columns["Visual Images Checked"], result.get("images_checked", "1"))

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def workbook_sheet_names(workbook_bytes: bytes) -> list[str]:
    workbook = openpyxl.load_workbook(io.BytesIO(workbook_bytes), read_only=True)
    return workbook.sheetnames


def workbook_row_count(workbook_bytes: bytes, sheet_name: str) -> int:
    workbook = openpyxl.load_workbook(io.BytesIO(workbook_bytes), read_only=True)
    return max(workbook[sheet_name].max_row - 1, 0)


def cache_counts(urls_by_row: list[tuple[int, str]], cache_path: str | Path) -> dict[str, int]:
    unique_urls = {url for _, url in urls_by_row}
    cache = load_cache(cache_path)
    return {
        "rows_with_urls": len(urls_by_row),
        "unique_urls": len(unique_urls),
        "remaining": sum(1 for url in unique_urls if url not in cache),
        "cached": sum(1 for url in unique_urls if url in cache),
        "retryable": sum(1 for url in unique_urls if cache.get(url, {}).get("verdict") in TRANSIENT_FAILURE_VERDICTS),
        "blocked": sum(1 for url in unique_urls if cache.get(url, {}).get("verdict") == "BLOCKED"),
        "not_found": sum(1 for url in unique_urls if cache.get(url, {}).get("verdict") == "NOT FOUND"),
        "negative": sum(1 for url in unique_urls if cache.get(url, {}).get("verdict") in {"NO", "NO IMAGE"}),
    }


def completed_results_by_row(urls_by_row: list[tuple[int, str]], cache_path: str | Path) -> dict[int, dict[str, str]]:
    cache = load_cache(cache_path)
    return {row: cache[url] for row, url in urls_by_row if url in cache}


def response_usage(response: Any) -> dict[str, int]:
    usage = getattr(response, "usage", None)
    if not usage:
        return zero_usage()
    input_tokens = int(usage_field(usage, "input_tokens") or 0)
    output_tokens = int(usage_field(usage, "output_tokens") or 0)
    details = usage_field(usage, "input_tokens_details")
    cached_tokens = int(usage_field(details, "cached_tokens") or 0) if details else 0
    return {"input_tokens": input_tokens, "cached_input_tokens": cached_tokens, "output_tokens": output_tokens}


def usage_field(value: Any, name: str) -> Any:
    if isinstance(value, dict):
        return value.get(name)
    return getattr(value, name, None)


def zero_usage() -> dict[str, int]:
    return {"input_tokens": 0, "cached_input_tokens": 0, "output_tokens": 0}


def add_usage(left: dict[str, int], right: dict[str, int]) -> dict[str, int]:
    return {
        "input_tokens": int(left.get("input_tokens", 0)) + int(right.get("input_tokens", 0)),
        "cached_input_tokens": int(left.get("cached_input_tokens", 0)) + int(right.get("cached_input_tokens", 0)),
        "output_tokens": int(left.get("output_tokens", 0)) + int(right.get("output_tokens", 0)),
    }


def estimate_cost(model: str, usage: dict[str, int]) -> float:
    pricing = MODEL_PRICING_PER_1M.get(model, MODEL_PRICING_PER_1M[DEFAULT_MODEL])
    cached_input = int(usage.get("cached_input_tokens", 0))
    total_input = int(usage.get("input_tokens", 0))
    uncached_input = max(total_input - cached_input, 0)
    output = int(usage.get("output_tokens", 0))
    return ((uncached_input * pricing["input"]) + (cached_input * pricing["cached_input"]) + (output * pricing["output"])) / 1_000_000


def label_fetch_failure(exc: Exception | None = None, reason: str = "") -> tuple[str, str]:
    reason = reason or (f"{type(exc).__name__}: {exc}" if exc else "")
    status_match = re.search(r"(\d{3}) Client Error", reason)
    status_code = status_match.group(1) if status_match else ""
    if status_code == "404":
        return "NOT FOUND", reason
    if status_code in {"401", "403", "406"}:
        return "BLOCKED", reason
    if status_code == "429":
        return "RATE LIMITED", reason
    if any(term in reason for term in ("Timeout", "ConnectionError", "SSLError")):
        return "FETCH ERROR", reason
    return "ERROR", reason


def normalize_cache_labels(cache: dict[str, dict[str, str]]) -> dict[str, dict[str, str]]:
    for result in cache.values():
        if result.get("verdict") != "ERROR":
            continue
        verdict, _ = label_fetch_failure(reason=result.get("reason", ""))
        result["verdict"] = verdict
    return cache


def image_result(candidate: ImageCandidate, verdict: str, reason: str, usage: dict[str, int], checked: int, skipped: int) -> dict[str, str]:
    return {
        "verdict": verdict,
        "image_url": candidate.url,
        "reason": reason,
        "source": candidate.source,
        "images_checked": str(checked),
        "images_skipped": str(skipped),
        "input_tokens": str(usage["input_tokens"]),
        "cached_input_tokens": str(usage["cached_input_tokens"]),
        "output_tokens": str(usage["output_tokens"]),
    }


def empty_result(verdict: str, reason: str) -> dict[str, str]:
    return {
        "verdict": verdict,
        "image_url": "",
        "reason": reason,
        "source": "",
        "images_checked": "0",
        "images_skipped": "0",
        "input_tokens": "0",
        "cached_input_tokens": "0",
        "output_tokens": "0",
    }


def dedupe_candidates(candidates: list[ImageCandidate]) -> list[ImageCandidate]:
    seen = set()
    deduped = []
    for candidate in candidates:
        if candidate.url in seen:
            continue
        seen.add(candidate.url)
        deduped.append(candidate)
    return deduped


def dedupe_pending(urls_by_row: list[tuple[int, str]]) -> list[tuple[int, str]]:
    seen = set()
    deduped = []
    for row, url in urls_by_row:
        if url in seen:
            continue
        seen.add(url)
        deduped.append((row, url))
    return deduped


def retry_pending(urls_by_row: list[tuple[int, str]], cache: dict[str, dict[str, str]], verdicts: set[str], retry_kind: str, batch_size: int) -> list[tuple[int, str]]:
    retry_count_key = "error_retry_count" if retry_kind == "error" else "negative_recheck_count"
    rows = []
    for row, url in dedupe_pending(urls_by_row):
        cached = cache.get(url, {})
        if cached.get("verdict") not in verdicts:
            continue
        rows.append((parse_int(cached.get(retry_count_key)), float_or_zero(cached.get("last_attempted_at")), row, url))
    rows.sort()
    return [(row, url) for _, _, row, url in rows[:batch_size]]


def with_retry_metadata(result: dict[str, str], previous: dict[str, str] | None, retry_kind: str | None) -> dict[str, str]:
    updated = dict(result)
    previous = previous or {}
    updated["last_attempted_at"] = f"{time.time():.6f}"
    updated["error_retry_count"] = str(parse_int(previous.get("error_retry_count")))
    updated["negative_recheck_count"] = str(parse_int(previous.get("negative_recheck_count")))
    if retry_kind == "error":
        updated["error_retry_count"] = str(parse_int(previous.get("error_retry_count")) + 1)
    if retry_kind == "negative":
        updated["negative_recheck_count"] = str(parse_int(previous.get("negative_recheck_count")) + 1)
    return updated


def result_usage(result: dict[str, str]) -> dict[str, int]:
    return {
        "input_tokens": parse_int(result.get("input_tokens")),
        "cached_input_tokens": parse_int(result.get("cached_input_tokens")),
        "output_tokens": parse_int(result.get("output_tokens")),
    }


def parse_int(value: Any) -> int:
    try:
        return int(str(value).strip())
    except (TypeError, ValueError):
        return 0


def float_or_zero(value: Any) -> float:
    try:
        return float(str(value).strip())
    except (TypeError, ValueError):
        return 0.0


def chunked(items: list[tuple[int, str]], size: int) -> list[list[tuple[int, str]]]:
    return [items[index : index + size] for index in range(0, len(items), size)]
